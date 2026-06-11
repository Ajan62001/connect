"""Office extraction (.xlsx/.xls/.docx): text shape + caps, routing
detection (content-type / extension / magic-byte tiebreaks), and the
pipeline end-to-end — no network (fixtures are generated in-test with
openpyxl/python-docx; xlrd cannot WRITE .xls, so a tiny xlwt-generated
binary fixture is committed at tests/fixtures/annex.xls instead)."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest

from canned_web import FakeFetcher, html_page
from connect.ingestion import extract_office as eo
from connect.ingestion.extract_html import ExtractionError
from connect.ingestion.extract_office import (
    detect_office_kind,
    extract_docx,
    extract_office,
    extract_xls,
    extract_xlsx,
)
from connect.storage import documents as doc_dao
from connect.storage import fts as fts_dao
from connect.storage import links as link_dao

XLS_FIXTURE = Path(__file__).parent / "fixtures" / "annex.xls"

XLSX_CTYPE = ("application/vnd.openxmlformats-officedocument"
              ".spreadsheetml.sheet")
XLS_CTYPE = "application/vnd.ms-excel"
DOCX_CTYPE = ("application/vnd.openxmlformats-officedocument"
              ".wordprocessingml.document")


# --- fixture builders (write to tmp_path, read bytes back) -------------------


def _xlsx_bytes(tmp_path, sheets: dict[str, list[list]]) -> bytes:
    """sheets: name -> rows (lists of cell values; None stays None)."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # drop the default sheet; build exactly what's asked
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        for row in rows:
            ws.append(row)
    path = tmp_path / "fixture.xlsx"
    wb.save(path)
    return path.read_bytes()


def _docx_bytes(tmp_path, paragraphs: list[str],
                table_rows: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_rows:
        table = document.add_table(rows=len(table_rows),
                                   cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, value in enumerate(row):
                table.rows[r].cells[c].text = value
    path = tmp_path / "fixture.docx"
    document.save(path)
    return path.read_bytes()


def _budget_xlsx(tmp_path) -> bytes:
    """The flagship shape: an RBI/budget-style annex, numbers verbatim."""
    return _xlsx_bytes(tmp_path, {
        "Statewise": [
            ["State", "Allocation", "Share"],
            ["Maharashtra", 182000, 0.18],
            [None, None, None],            # fully empty -> skipped
            ["Karnataka", 95000, 1.85],
        ],
        "Notes": [["Figures in lakh rupees"]],
    })


# --- xlsx ---------------------------------------------------------------------


def test_extract_xlsx_sheets_rows_and_numbers(tmp_path):
    extracted = extract_xlsx(_budget_xlsx(tmp_path))
    text = extracted.text

    assert "## Statewise" in text
    assert "## Notes" in text
    assert "State\tAllocation\tShare" in text
    assert "Maharashtra\t182000\t0.18" in text   # numbers verbatim
    assert "Karnataka\t95000\t1.85" in text
    assert "182000.0" not in text
    # the fully-empty row was skipped (no blank line between data rows)
    statewise = text.split("## Notes")[0]
    assert "\n\n" not in statewise.strip()
    assert "[truncated" not in text              # no caps hit


def test_extract_xlsx_sheet_cap(tmp_path):
    data = _xlsx_bytes(tmp_path, {
        f"Sheet{i:02d}": [[f"row in sheet {i}"]] for i in range(12)})
    text = extract_xlsx(data).text
    assert "## Sheet09" in text
    assert "## Sheet10" not in text
    assert "[truncated: only first 10 of 12 sheets extracted]" in text


def test_extract_xlsx_row_cap(tmp_path):
    rows = [[f"item {i}", i] for i in range(505)]
    text = extract_xlsx(_xlsx_bytes(tmp_path, {"Data": rows})).text
    assert "item 499" in text
    assert "item 500" not in text
    assert "[truncated: only first 500 rows of sheet 'Data' extracted]" in text


def test_extract_xlsx_column_cap(tmp_path):
    wide = [[f"c{i}" for i in range(55)]]
    text = extract_xlsx(_xlsx_bytes(tmp_path, {"Wide": wide})).text
    assert "c49" in text
    assert "c50" not in text
    assert ("[truncated: only first 50 columns of sheet 'Wide' extracted]"
            in text)


def test_extract_xlsx_empty_workbook_raises(tmp_path):
    data = _xlsx_bytes(tmp_path, {"Blank": [[None, None]]})
    with pytest.raises(ExtractionError):
        extract_xlsx(data)


def test_extract_xlsx_garbage_raises():
    with pytest.raises(ExtractionError):
        extract_xlsx(b"PK\x03\x04 this is not a real zip")


# --- xls (committed binary fixture: xlrd cannot write .xls) ---------------------


def test_extract_xls_fixture():
    data = XLS_FIXTURE.read_bytes()
    assert data[:4] == b"\xd0\xcf\x11\xe0"  # OLE magic — the routing tiebreak
    extracted = extract_xls(data)
    text = extracted.text

    assert "## Summary" in text
    assert "## Notes" in text
    assert "State\tAllocation" in text
    # XLS stores numbers as floats; integral values must come out verbatim
    assert "Maharashtra\t182000" in text
    assert "182000.0" not in text
    assert "Karnataka\t95000.5" in text
    assert "Figures in lakh rupees" in text
    # the fully-empty row 3 was skipped
    assert "Karnataka\t95000.5\nTotal\t277000.5" in text


def test_extract_xls_garbage_raises():
    with pytest.raises(ExtractionError):
        extract_xls(b"\xd0\xcf\x11\xe0 truncated OLE junk")


# --- docx ----------------------------------------------------------------------


def test_extract_docx_paragraphs_and_table(tmp_path):
    data = _docx_bytes(
        tmp_path,
        paragraphs=["Committee Report", "",  # empty paragraph dropped
                    "The committee recommends revising the threshold."],
        table_rows=[["State", "Grant"], ["Maharashtra", "182000"]])
    extracted = extract_docx(data)
    text = extracted.text

    assert "Committee Report" in text
    assert "The committee recommends revising the threshold." in text
    assert "State\tGrant" in text
    assert "Maharashtra\t182000" in text


def test_extract_docx_table_row_cap(tmp_path, monkeypatch):
    monkeypatch.setattr(eo, "MAX_ROWS", 5)  # same cap logic, kept fast
    rows = [[f"row {i}", "x"] for i in range(8)]
    data = _docx_bytes(tmp_path, paragraphs=["Annex"], table_rows=rows)
    text = extract_docx(data).text
    assert "row 4" in text
    assert "row 5" not in text
    assert "[truncated: only first 5 rows of table 1 extracted]" in text


def test_extract_docx_garbage_raises():
    with pytest.raises(ExtractionError):
        extract_docx(b"PK\x03\x04 nope")


# --- routing detection ------------------------------------------------------------


def test_detect_by_extension(tmp_path):
    xlsx = _budget_xlsx(tmp_path)
    assert detect_office_kind(
        content_type="application/octet-stream",
        name="https://rbi.org.in/rdocs/Annex1.XLSX", data=xlsx) == "xlsx"
    assert detect_office_kind(
        content_type="", name="budget.docx",
        data=_docx_bytes(tmp_path, ["x"])) == "docx"
    assert detect_office_kind(
        content_type=None, name="annex.xls",
        data=XLS_FIXTURE.read_bytes()) == "xls"
    # query strings never count as the extension
    assert detect_office_kind(
        content_type="text/html", name="https://x.gov.in/page?file=a.xlsx",
        data=b"<html>") is None


def test_detect_by_content_type(tmp_path):
    xlsx = _budget_xlsx(tmp_path)
    # extension-less gov download endpoints — header decides
    assert detect_office_kind(
        content_type=XLSX_CTYPE,
        name="https://rbi.org.in/Scripts/Download.aspx?id=9",
        data=xlsx) == "xlsx"
    assert detect_office_kind(
        content_type=DOCX_CTYPE, name="upload",
        data=_docx_bytes(tmp_path, ["x"])) == "docx"
    assert detect_office_kind(
        content_type=XLS_CTYPE, name="download",
        data=XLS_FIXTURE.read_bytes()) == "xls"


def test_detect_magic_tiebreaks(tmp_path):
    xlsx = _budget_xlsx(tmp_path)
    ole = XLS_FIXTURE.read_bytes()
    # modern zip payload behind a legacy label -> xlsx
    assert detect_office_kind(
        content_type=XLS_CTYPE, name="annex.xls", data=xlsx) == "xlsx"
    # legacy OLE payload behind a modern label -> xls
    assert detect_office_kind(
        content_type=XLSX_CTYPE, name="annex.xlsx", data=ole) == "xls"
    # magic alone never initiates a match: a random zip is NOT a spreadsheet
    assert detect_office_kind(
        content_type="application/zip", name="bundle.zip", data=xlsx) is None
    assert detect_office_kind(
        content_type="", name="https://x.gov.in/thing", data=ole) is None
    # and PDFs / HTML stay untouched
    assert detect_office_kind(
        content_type="application/pdf", name="a.pdf", data=b"%PDF") is None
    assert detect_office_kind(
        content_type="text/html", name="https://x.gov.in/a",
        data=b"<html>") is None


def test_extract_office_zip_sibling_fallback(tmp_path):
    """A .docx served with an .xlsx label: openpyxl fails, the sibling OOXML
    extractor recovers (both are 'PK' zips — labels lie on gov servers)."""
    docx_data = _docx_bytes(tmp_path, ["The recovered paragraph."])
    extracted = extract_office(docx_data, "xlsx")
    assert "The recovered paragraph." in extracted.text


# --- pipeline end-to-end (no network: upload + canned fetcher) -------------------


async def test_ingest_file_xlsx_end_to_end(tmp_path, container, db):
    data = _budget_xlsx(tmp_path)
    result = await container.pipeline.ingest_file(
        db, "rbi-annex.xlsx", data, XLSX_CTYPE)

    assert result.created is True
    doc = result.document
    assert doc.media_type == "xlsx"
    assert doc.title == "rbi-annex.xlsx"  # filename fallback, like PDFs
    assert "Maharashtra\t182000\t0.18" in doc.content_text

    # raw bytes are blob-stored and the row is FTS-searchable
    from dbutil import qv
    blob_path = await qv(db, "SELECT raw_blob_path FROM document"
                             " WHERE id=%s", doc.id)
    assert container.blobs.get(blob_path) == data
    items, total = await fts_dao.search_documents(db, "Maharashtra")
    assert total == 1 and items[0].id == doc.id
    assert items[0].media_type == "xlsx"


async def test_ingest_file_xls_and_docx(tmp_path, container, db):
    xls = await container.pipeline.ingest_file(
        db, "annex.xls", XLS_FIXTURE.read_bytes(), XLS_CTYPE)
    assert xls.document.media_type == "xlsx"  # one spreadsheet bucket
    assert "Karnataka\t95000.5" in xls.document.content_text

    docx = await container.pipeline.ingest_file(
        db, "report.docx",
        _docx_bytes(tmp_path, ["A standalone committee report draft."]),
        DOCX_CTYPE)
    assert docx.document.media_type == "docx"
    items, total = await fts_dao.search_documents(db, "standalone")
    assert total == 1 and items[0].id == docx.document.id


async def test_ingest_file_bad_office_takes_failure_path(container, db):
    """A corrupt spreadsheet fails exactly like a bad PDF: ExtractionError,
    no document row created."""
    from dbutil import qv
    with pytest.raises(ExtractionError):
        await container.pipeline.ingest_file(
            db, "broken.xlsx", b"PK\x03\x04 not really a workbook",
            XLSX_CTYPE)
    assert await qv(db, "SELECT COUNT(*) FROM document") == 0


async def test_ingest_url_xlsx_by_content_type(tmp_path, container, db):
    url = "https://www.rbi.org.in/Scripts/AnnexDownload.aspx?ID=42"
    container.pipeline.fetcher = FakeFetcher(
        {url: (XLSX_CTYPE, _budget_xlsx(tmp_path))})

    result = await container.pipeline.ingest_url(db, url, title="Annex I")
    assert result.document.media_type == "xlsx"
    assert result.document.title == "Annex I"
    assert "Karnataka\t95000\t1.85" in result.document.content_text


async def test_link_follow_fetches_xlsx_annex(tmp_path, container, db):
    """The motivating case: an RBI page links its .xlsx annex (served with a
    lying legacy content-type); the follower ingests it as a document."""
    parent = "https://www.rbi.org.in/Scripts/NotificationUser.aspx?Id=12999"
    annex = "https://www.rbi.org.in/rdocs/content/xls/Annex12999.xlsx"
    container.pipeline.fetcher = FakeFetcher({
        parent: html_page(
            "Circular with a data annex",
            f'<p><a href="{annex}">State-wise data annex</a></p>'),
        annex: (XLS_CTYPE, _budget_xlsx(tmp_path)),  # PK magic wins -> xlsx
    })

    result = await container.pipeline.ingest_url(db, parent)
    links = await link_dao.list_for_document(db, result.document.id)
    assert len(links) == 1 and links[0].status == "fetched"

    child = await doc_dao.get(db, links[0].resolved_document_id)
    assert child is not None
    assert child.media_type == "xlsx"
    assert "Maharashtra\t182000\t0.18" in child.content_text


async def test_link_follow_bad_xlsx_marks_link_failed(container, db):
    parent = "https://www.rbi.org.in/Scripts/NotificationUser.aspx?Id=13000"
    annex = "https://www.rbi.org.in/rdocs/content/xls/Broken.xlsx"
    container.pipeline.fetcher = FakeFetcher({
        parent: html_page(
            "Circular with a broken annex",
            f'<p><a href="{annex}">data annex</a></p>'),
        annex: (XLSX_CTYPE, b"PK\x03\x04 broken workbook bytes"),
    })

    result = await container.pipeline.ingest_url(db, parent)  # no raise
    link = (await link_dao.list_for_document(db, result.document.id))[0]
    assert link.status == "failed"
    assert link.resolved_document_id is None
